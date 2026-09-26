"""Generate the editable Word CV from the authoritative standalone LaTeX CV.

Requires python-docx. Supports the specific macros used by this CV, not arbitrary TeX.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'Muhammad-Tahir-CV-2026-Sept.tex'
OUTPUT = SOURCE.with_suffix('.docx')

def argument(text, pos):
    while text[pos].isspace(): pos += 1
    assert text[pos] == '{', text[pos:pos+100]
    start = pos + 1; depth = 1; pos += 1
    while depth:
        if text[pos] == '{' and text[pos-1] != '\\': depth += 1
        if text[pos] == '}' and text[pos-1] != '\\': depth -= 1
        pos += 1
    return text[start:pos-1], pos

def plain(text):
    text = re.sub(r'\\(?:Needspace|needspace|vspace)\{[^}]*\}', '', text)
    text = re.sub(r'\\(?:begin|end)\{itemize\}', '', text)
    text = re.sub(r'\\(?:ifshowupworkavailability|fi|par|nopagebreak)\b', '', text)
    text = re.sub(r'\\\\(?:\[[^]]*\])?', ' ', text)
    pattern = re.compile(r'\\(textbf|textit|mbox|href|url)\s*')
    while match := pattern.search(text):
        value, end = argument(text, match.end())
        if match.group(1) == 'href': value, end = argument(text, end)
        text = text[:match.start()] + value + text[end:]
    text = text.replace('\\&', '&').replace('~', ' ').replace('---', '-').replace('--', '-')
    text = text.replace('``', '"').replace("''", '"')
    return re.sub(r'\s+', ' ', text).strip()

def hyperlink(paragraph, label, url):
    element = OxmlElement('w:hyperlink')
    element.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '000000' if paragraph.style.name.startswith('Heading') else '00467F')
    props.append(color)
    run.append(props); text = OxmlElement('w:t'); text.text = label; run.append(text)
    element.append(run); paragraph._p.append(element)

def rich(paragraph, source):
    while match := re.search(r'\\(?:textbf|textit|mbox)\s*', source):
        value, end = argument(source, match.end())
        source = source[:match.start()] + value + source[end:]
    pos = 0
    for match in re.finditer(r'\\href\s*', source):
        if match.start() < pos: continue
        url, end = argument(source, match.end()); label, end = argument(source, end)
        before = plain(source[pos:match.start()])
        if before: paragraph.add_run(before + ' ')
        hyperlink(paragraph, plain(label), url)
        pos = end
    remaining = plain(source[pos:])
    if remaining: paragraph.add_run((' ' if pos and remaining[0] not in '.,;:!?' else '') + remaining)

def build():
    source = re.sub(r'(?<!\\)%[^\n]*', '', SOURCE.read_text(encoding='utf-8-sig'))
    body = source[source.index('\\section{Professional Summary}'):source.index('\\end{document}')]
    doc = Document(); sec = doc.sections[0]
    sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.3)
    sec.left_margin = sec.right_margin = Cm(1.4)
    sec.footer_distance = Cm(.55)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet']:
        style = doc.styles[name]; style.font.name = 'Arial'; style.font.color.rgb = RGBColor(0,0,0)
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(1)
        style.paragraph_format.line_spacing = 1.0
    normal = doc.styles['Normal'].paragraph_format
    normal.widow_control = True
    for name in ['Heading 1', 'Heading 2']:
        doc.styles[name].paragraph_format.keep_with_next = True
        doc.styles[name].paragraph_format.space_before = Pt(5)
        doc.styles[name].font.bold = True
    doc.styles['Heading 1'].font.size = Pt(12)
    doc.styles['Title'].font.size = Pt(20); doc.styles['Title'].font.bold = True
    for border in list(doc.styles['Title'].element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    doc.styles['Footer'].font.name = 'Arial'
    doc.styles['Footer'].font.size = Pt(8)
    doc.styles['Footer'].font.color.rgb = RGBColor(100,100,100)
    bullets = doc.styles['List Bullet'].paragraph_format
    bullets.left_indent = Cm(.45); bullets.first_line_indent = Cm(-.3)
    p = doc.add_paragraph('MUHAMMAD TAHIR KOREJO', 'Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('+39-3317606955 | '); hyperlink(p,'tahirkorejo872@gmail.com','mailto:tahirkorejo872@gmail.com'); p.add_run(' | Genova, Italy')
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('LinkedIn: '); hyperlink(p,'muhammad-tahir','https://www.linkedin.com/in/muhammad-tahir-83a431204/')
    p.add_run(' | Portfolio: '); hyperlink(p,'muhammadtahir31.github.io/cv/','https://muhammadtahir31.github.io/cv/')
    p.paragraph_format.space_after = Pt(5)
    token = re.compile(r'\\(section|cventry|cvproject|item\b|newpage\b|begin\{tabularx\})')
    pos = 0
    def prose(chunk):
        if '\\par' in chunk:
            for part in re.split(r'\\par\b', chunk): prose(part)
            return
        if plain(chunk):
            p = doc.add_paragraph(); rich(p, chunk)
            if chunk.strip().startswith('\\textbf'):
                for run in p.runs: run.bold = True
                p.paragraph_format.keep_with_next = True
    while match := token.search(body, pos):
        prose(body[pos:match.start()]); kind = match.group(1); pos = match.end()
        if kind == 'newpage':
            doc.add_page_break()
        elif kind == 'section':
            title,pos = argument(body,pos); doc.add_paragraph(plain(title),'Heading 1')
        elif kind in ['cventry','cvproject']:
            args=[]
            for _ in range(3 if kind=='cventry' else 2):
                value,pos=argument(body,pos);args.append(value)
            label,date = (args[1],args[2]) if kind=='cventry' else args
            p=doc.add_paragraph(style='Heading 2')
            p.paragraph_format.tab_stops.add_tab_stop(Cm(18.2),WD_TAB_ALIGNMENT.RIGHT)
            rich(p,label);p.add_run('\t'+plain(date))
            if kind=='cventry':
                p=doc.add_paragraph(plain(args[0]));p.paragraph_format.keep_with_next=True
                for run in p.runs:run.italic=True
        elif kind=='item':
            nxt=token.search(body,pos);end=nxt.start() if nxt else len(body)
            chunk=body[pos:end]
            # End the bullet before any text outside its itemize environment.
            boundary=chunk.find('\\end{itemize}')
            if boundary>=0:end=pos+boundary
            p=doc.add_paragraph(style='List Bullet');rich(p,body[pos:end]);pos=end
        else:
            _,pos=argument(body,pos);_,pos=argument(body,pos)
            end=body.index('\\end{tabularx}',pos)
            for row in re.split(r'\\\\(?:\[[^]]*\])?',body[pos:end]):
                if '&' not in row:continue
                label,value=re.split(r'(?<!\\)&',row,maxsplit=1)
                p=doc.add_paragraph();p.add_run(plain(label)+' ').bold=True;p.add_run(plain(value))
            pos=end+len('\\end{tabularx}')
    prose(body[pos:])
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Muhammad Tahir Korejo | Page ')
    for code,tail in [('PAGE',' of '),('NUMPAGES','')]:
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),code);footer._p.append(field)
        if tail:footer.add_run(tail)
    for run in footer.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor(100,100,100)
    doc.core_properties.title='Muhammad Tahir Korejo CV'
    doc.core_properties.author='Muhammad Tahir Korejo'
    doc.core_properties.subject='Java and Spring Boot software engineering'
    doc.save(OUTPUT)
    print(f'Created {OUTPUT.name} from {SOURCE.name}')

if __name__=='__main__': build()
